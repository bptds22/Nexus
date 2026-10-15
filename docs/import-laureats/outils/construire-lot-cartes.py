"""Import Lauréats — construit le lot des cartes prospect depuis la liste de dépistage Excel.

Usage (depuis un dossier sans code, par exemple le scratchpad) :
    python construire-lot-cartes.py <fichier.xlsx> <dossier_sortie>

Le dossier de sortie DOIT être hors git (data/coachs/… est ignoré) : le lot et l'aperçu
contiennent les noms et les notes des joueurs. Ce script ne contient aucune donnée personnelle.

Sorties :
  lot-cartes-laureats.sql   — transaction gardée, UN SEUL envoi (décision BP 2026-10-15)
  apercu-cartes-laureats.txt — compte, grades, équipes, notes « à confirmer », doublons, extractions
  ids-cartes-laureats.txt   — les id des cartes (servent au rollback)
Les id sont déterministes (uuid5 sur onglet + ligne Excel) : relancer le script redonne les mêmes.
"""
import sys, re, uuid, datetime, unicodedata, collections
import openpyxl

ESPACE = uuid.UUID("5b0f7a52-9a51-4c43-8f0c-6c1d3a0e2b27")  # fixe : rend les id reproductibles
UNITE_CEGEP = "dfad9036-0d54-4304-906f-91f19542a882"           # Cégep de Saint-Hyacinthe
PROMOTION = 2027
MD5_AVANT_INSERT = "10e4d1ea3593029540e460e25b597823"
TRIGGERS = ["trg_carte_a_avant_insert", "trg_carte_b_avant_update", "trg_carte_c_fusion", "trg_carte_c_invitation",
            "trg_carte_d_telephone", "trg_carte_z_inviter", "trg_carte_z_inviter_ajout", "trg_carte_z_journal",
            "trg_carte_z_suppression", "trg_carte_z_telephone", "trg_rapprochement_carte_insert",
            "trg_rapprochement_carte_update"]

# Onglet → équipe 2026-2027 (Juvénile ; Midget AAA pour les clubs civils). Relevé prod 2026-10-15.
EQUIPES = {
    "Académie Les Estacades": "b23b46e1-3bbb-4488-8b86-f3c91f8cba83",
    "Académie Saint-Louis": "9a7ee739-4c66-40c4-a86b-eff77a4ca16d",
    "Amitié": "28b005a9-d83f-4d67-bb38-8302e0201e81",
    "Baie-St-François": "68d0c6e1-8e6c-4c13-9060-84584570c4ed",
    "Boucherville": "c46a1a7e-f937-4229-a2f6-dee6b07d161d",
    "Bourget": "a5faab0b-7c8a-415b-a5da-48f92ce99f83",
    "Cité-Des-Jeunes": "4b41e4db-8999-4ff1-8178-bc8bf1c378c0",
    "Col. Charles-Lemoyne": "66f527f9-dff7-400a-8e9a-2ed0eb845fc5",
    "Col. De Lévis": "f7fe9a38-7192-4081-8744-fbfb238aadff",
    "Col. De Montréal": "9c229d71-0340-4d29-aada-13cd049667eb",
    "Col. Mont-Sacré-Coeur": "62d6887b-4fd6-4a64-b7db-66c6afdc9f86",
    "Col. Saint-Jean-Vianney": "a03beb90-ca2d-4fdc-b6b9-d0c46980591d",
    "Courvilloise": "b886a6ae-2b64-4ac1-ab3d-c98097b96f70",
    "Curé-Antoine-Labelle": "2cbb4aa6-8a8d-4c98-bfa0-97c9653992fa",
    "Donnacona": "4030f4e7-1af3-4b29-bce6-6aa9362686b7",
    "Du Phare": "5f12474f-14f2-4afe-91e5-d317dc9626c4",
    "Greenfield Park": "60a656d9-573f-487a-8325-69faefa0dff1",
    "J.-H.-Leclerc": "639e23cf-dcb2-4dff-a049-855aa99af302",
    "Jacques Rousseau": "3b17e5eb-8989-4ab4-90f6-4d0af6b280c1",
    "Jean-Eudes": "5dac9a66-e503-4fe2-95f5-80ee5b027664",
    "La Frontalière": "1cb4067e-bec0-4d45-bcba-b35ed471c9d0",
    "La Ruche": "1eda9693-9460-45e5-9986-656e046f8a86",
    "Laprairie": "1a340d5d-1090-460d-ba0d-c7ae091f9073",
    "LAssomption": "6ad77d53-cebc-4a81-8492-a871157cc4d0",
    "Le Ber": "893898b4-cda8-4681-8865-8b26ca4d6d57",
    "Letendre": "eb0d11c3-908e-4414-a88b-2ec309c6e754",
    "Montcalm": "525725a2-d9bc-44b9-8978-e108918676ee",
    "Mt St-Sacrement": "e67a9a05-73f3-47f1-b4e2-f16eeb284905",
    "Nicolas-Gatineau": "1edd79c4-d7b6-4c9c-b294-f8f12069ea2b",
    "Notre-Dame": "d40ffddf-1019-4389-a278-d2ca32dee0b1",
    "Nouvelles-Frontières": "348df7f0-2b01-4c95-b42e-a41fe70a3884",
    "PGO": "9039a210-63dd-4104-af55-23ccd7c676a5",
    "Pointe-Lévy": "b5f8cf79-9e4e-465b-baf0-135d8d70051a",
    "Poly. Bélanger": "87fa17c7-0103-4c94-a086-cc6a924b5c59",
    "Poly. Charlesbourg": "1274cf05-1f77-4206-b9a7-406695243b14",
    "Poly. Deux-Montagnes": "81279b35-f4e1-4a1c-b684-db616664f724",
    "Poly. Montignac": "2003073e-c264-47e8-8b8b-6042d9d07991",
    "Poly. St-Francois": "cabde120-3e5b-47c6-aac0-d4282f07a3d4",
    "Richelieu": "a160a230-fe01-4788-915c-e9230b16783e",
    "Roger-Comtois": "aea474f6-d5a5-4239-a495-093ed43921bb",
    "Saint-Jean-Eudes": "ffcad021-cb9b-4286-bdfb-7528494040e0",
    "Saint-Joseph STH": "c1d840d9-7af5-446d-835a-6470a793f163",
    "Sém. Saint-François": "12e706b6-febd-47cb-81b6-f18ffebfda89",
    "Sém. Saint-Joseph": "b2c0093d-d231-41dc-9a62-b9b1e3df847c",
    "Sentiers": "9791b531-8979-4829-b5a5-a81b5c25adf1",
    "St-Charles Garnier": "a9c7bde7-9225-4757-b322-7115be751f22",
    "St-Patricks": "0836bbd7-b294-4d31-8f7d-3daee7588b81",
    "Thérèse-Martin": "5e22f235-407e-47cb-ad8e-febfd28281e9",
    "Triolet HARFANGS": "3a68b77a-44fd-492a-98cc-b4595bebaa8d",
    "Victo Vicas": "8f755fdd-36b8-4d25-920f-949ff6e1f167",
    "Wilbrod-Dufour": "c3e58603-178e-49e0-9703-53682c8588bc",
}
NOM_PRENOM = {"Académie Les Estacades", "Bourget", "Col. Charles-Lemoyne", "Col. De Lévis", "Col. De Montréal",
              "Curé-Antoine-Labelle", "Donnacona", "J.-H.-Leclerc", "Richelieu"}

# Positions football (abréviation → id), relevé prod 2026-10-15.
POSITIONS = {
    "C": "e28f9ea7-2940-4ee6-af1d-43eea1b09523", "CB": "0ef3b923-3625-4e39-8323-0f35644f9dac",
    "DE": "7dc94660-5dad-4cd9-a481-a6aae9b6ff4c", "DL": "245d1f9a-09ee-425f-a598-bc0a899155b8",
    "DT": "57b4e169-247d-428a-98f8-854877f44a92", "FB": "11d9472f-19eb-4707-8f5d-344182da80d2",
    "FS": "823b8d75-3905-413b-8624-a3f7706f4690", "ILB": "13057eb2-3f3a-45bd-b2ee-3d02cbf6aad4",
    "K": "26557fd0-5aaa-42f5-9ffe-dc71a3e0aee3", "LB": "fb8f8e46-026a-48dc-839b-f0b6bb18fb0b",
    "LS": "16610770-f88b-4db0-9ccb-8db0a6c38e50", "OG": "4fa52bad-d26e-4eea-b0e7-ce6100729d9b",
    "OL": "f7d17505-e124-4dd0-886e-ff3605ce721e", "OLB": "2f0612cc-8f15-4674-bb7e-87fce536051c",
    "OT": "59a39573-e640-444d-a894-c2430d3029de", "P": "fa0314d5-cec9-4237-bdc6-08abce5871f1",
    "QB": "0d8ed2ce-a0c6-4d06-a921-3aef5dda94f6", "RB": "a0127a71-4cb7-408f-a765-7e37397c3bc3",
    "RET": "2e207a64-9d56-43e4-9e05-58cca36e2b9a", "S": "a3e612c3-9e98-4346-b128-f1da8848fe48",
    "SS": "16134eec-9e40-4a99-8b2b-b6f59e2c6db5", "TE": "333b9a6d-16fc-4060-ac1b-327f65f94662",
    "WR": "dea59633-0813-47c8-b652-b583b9c5f083",
}
ALIAS_POSITION = {"REC": "WR", "WR": "WR", "SB": "WR", "LD": "DL", "DL": "DL", "DLINE": "DL", "LO": "OL", "OL": "OL",
                  "OLINE": "OL", "SEC": "LB", "LB": "LB", "QA": "QB", "QB": "QB", "HB": "RB", "DO": "RB", "RB": "RB",
                  "EDGE": "DE", "DE": "DE", "MLB": "ILB"}
SANS_POSITION = {"DB", "DD", "ATH", "DEMIS", "RC", "F"}
GRADES = {1.0: "A+", 1.5: "A", 2.0: "B+", 2.5: "B", 3.0: "C+", 3.5: "C", 4.0: "D+", 3.99: "D+", 4.5: "D", 5.0: "E"}
NOTE_NOM = "Correspondance à confirmer : nom/prénom"
NOTE_NOM_INCOMPLET = "Correspondance à confirmer : nom/prénom incomplet"
A_CONFIRMER = "(à confirmer)"
# Noms d'un seul mot — place du mot tranchée par BP le 2026-10-15 ; tout autre cas arrête le script.
SEUL_MOT = {("Boucherville", "Isaac"): "prenom", ("Triolet HARFANGS", "Sergio"): "prenom",
            ("Boucherville", "Boivin"): "nom"}
PARTICULES = {"de", "du", "des", "da", "di", "del", "della", "van", "von", "le", "la", "el", "al", "mc", "st", "saint"}


def texte(v):
    return "" if v is None else str(v).strip()


def cle_nom(s):
    s = unicodedata.normalize("NFKD", s.casefold())
    return " ".join("".join(c for c in s if not unicodedata.combining(c)).split())


def position(brut):
    """→ (abréviation ou None, note ou None). La 1re d'une saisie multiple."""
    raw = texte(brut)
    if not raw:
        return None, None
    premier = re.split(r"[-/\s(]", raw)[0].strip().rstrip("?").upper()
    if premier in SANS_POSITION:
        return None, f"Position Excel : {raw}"
    abr = ALIAS_POSITION.get(premier, premier if premier in POSITIONS else None)
    if abr is None:
        return None, f"Position Excel : {raw}"
    # Option (b), décision BP 2026-10-15 : un simple synonyme bien rattaché (REC→WR, LO→OL, qb→QB…) ne laisse
    # pas de note. La note reste pour une saisie multiple (« OL-DL », « REC / DB », « LB (OLB) ») ou un « ? ».
    multiple = re.search(r"[-/\s(]", raw) is not None
    return abr, (f"Position Excel : {raw}" if multiple or "?" in raw else None)


def numero(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return str(int(v)) if float(v).is_integer() and 0 <= v < 10000 else None
    s = texte(v)
    m = re.fullmatch(r"(\d{1,4})(?:\s+\D.*)?", s)
    return m.group(1) if m else None   # « 13 t'es qui? » → 13 ; « 8-16 », « x », vide → rien


def grade(cell):
    """→ (grade ou None, note ou None, valeur numérique lue)."""
    v = cell.value
    if v is None or texte(v) == "":
        return None, None, None
    if isinstance(v, datetime.datetime):            # format « d.m » : 05-03 s'affiche 3.5
        val = float(f"{v.day}.{v.month}")
    elif isinstance(v, (int, float)):
        val = float(v)
    else:
        return None, f"Rang Excel : {texte(v)}", None
    if val == 0:
        return None, "Rang Excel : 0 (à réévaluer)", val
    g = GRADES.get(round(val, 2))
    return (g, None, val) if g else (None, f"Rang Excel : {texte(v)}", val)


TAILLE = re.compile(r"(?<!\d)([4-7])\s*(?:'|’|′|pieds?\b|pi\b)\s*,?\s*(?:(\d{1,2})(?!\d)\s*(?:\"|″|”|''|pouces?\b|po\b)?)?", re.I)
POIDS_LBS = re.compile(r"(?<!\d)(\d{2,3})\s*(?:lbs?|livres)\b", re.I)
POIDS_APRES = re.compile(r"\s*[,/-]?\s*(\d{2,3})(?!\d)(?!\s*(?:'|’|pieds?|ans|%))")


def taille_poids(notes):
    """→ (pieds, pouces, poids, extrait) — rien si hors bornes (80–350 lb, pouces ≤ 11)."""
    pieds = pouces = poids = None
    m = TAILLE.search(notes)
    fin = None
    if m:
        p, po = int(m.group(1)), int(m.group(2)) if m.group(2) else 0
        if po <= 11:
            pieds, pouces, fin = p, po, m.end()
    for w in POIDS_LBS.finditer(notes):
        if re.search(r"perdu\s*$", notes[max(0, w.start() - 12):w.start()], re.I):
            continue                                # « a perdu 50 lbs » n'est pas un poids
        if 80 <= int(w.group(1)) <= 350:
            poids = int(w.group(1))
            break
    if poids is None and fin is not None:
        w = POIDS_APRES.match(notes, fin)
        if w and 80 <= int(w.group(1)) <= 350:
            poids = int(w.group(1))
    extrait = notes[m.start():m.end() + 12] if m else None
    return pieds, pouces, poids, extrait


def sql(v):
    if v is None:
        return "null"
    if isinstance(v, int):
        return str(v)
    return "'" + str(v).replace("'", "''") + "'"


def lire(chemin):
    wb = openpyxl.load_workbook(chemin, data_only=True)
    cartes, doublons, ignores, sans_nom = [], [], [], 0
    vus = {}
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            ignores.append(ws.title)
            continue
        entete = [texte(c.value).upper() for c in ws[10]][:6]
        if entete != ["ÉQUIPE", "NOM", "NUMERO", "POS", "RANK", "NOTES"]:
            raise SystemExit(f"en-tête inattendu dans {ws.title!r} : {entete}")
        for row in ws.iter_rows(min_row=11, max_col=6):
            valeurs = [c.value for c in row]
            if all(texte(v) == "" for v in valeurs):
                continue
            nom_complet = " ".join(texte(row[1].value).split())
            if not nom_complet:
                sans_nom += 1
                continue
            if ws.title not in EQUIPES:
                raise SystemExit(f"onglet sans équipe : {ws.title!r}")
            cle = (EQUIPES[ws.title], cle_nom(nom_complet))
            if cle in vus:
                doublons.append((ws.title, row[0].row, nom_complet, vus[cle]))
                continue
            vus[cle] = row[0].row
            mots = nom_complet.split(" ")
            notes = []
            if "," in nom_complet:                  # « Nom, Prénom » : la virgule tranche, quel que soit l'onglet
                nom, prenom = (" ".join(x.split()) for x in nom_complet.split(",", 1))
            elif len(mots) == 1:
                place = SEUL_MOT.get((ws.title, mots[0]))
                if place is None:
                    raise SystemExit(f"nom d'un seul mot non tranché par BP : {ws.title!r} L{row[0].row}")
                prenom, nom = (mots[0], A_CONFIRMER) if place == "prenom" else (A_CONFIRMER, mots[0])
                notes.append(NOTE_NOM_INCOMPLET)
            elif ws.title in NOM_PRENOM:
                k = 1
                while k < len(mots) - 1 and mots[k - 1].casefold() in PARTICULES:
                    k += 1                          # « De Lafontaine Victor » : nom « De Lafontaine »
                nom, prenom = " ".join(mots[:k]), " ".join(mots[k:])
            else:
                prenom, nom = mots[0], " ".join(mots[1:])
            if len(mots) > 2:
                notes.append(NOTE_NOM)
            abr, note_pos = position(row[3].value)
            g, note_rang, _ = grade(row[4])
            texte_notes = texte(row[5].value)
            pieds, pouces, poids, extrait = taille_poids(texte_notes) if texte_notes else (None, None, None, None)
            if texte_notes:
                notes.insert(0, texte_notes)
            notes += [n for n in (note_pos, note_rang) if n]
            cartes.append(dict(
                id=str(uuid.uuid5(ESPACE, f"{ws.title}|{row[0].row}")), onglet=ws.title, ligne=row[0].row,
                team_id=EQUIPES[ws.title], prenom=prenom, nom=nom, numero=numero(row[2].value),
                position=abr, position_id=POSITIONS.get(abr) if abr else None, pos_brute=texte(row[3].value),
                grade=g, rang_brut=row[4].value, pieds=pieds, pouces=pouces, poids=poids, extrait=extrait,
                notes=notes, nom_complet=nom_complet))
    return cartes, doublons, ignores, sans_nom


def ecrire_lot(cartes, chemin):
    n = len(cartes)
    nb_notes = sum(len(c["notes"]) for c in cartes)
    ids = ",\n ".join(f"('{c['id']}'::uuid)" for c in cartes)
    equipes = sorted({c["team_id"] for c in cartes})
    lignes = ",\n ".join(
        "(" + ", ".join([sql(c["id"]) + "::uuid", sql(c["team_id"]) + "::uuid", sql(c["prenom"]), sql(c["nom"]),
                         sql(c["numero"]), (sql(c["position_id"]) + "::uuid") if c["position_id"] else "null::uuid",
                         str(PROMOTION), sql(c["pieds"]), sql(c["pouces"]), sql(c["poids"]), sql(c["grade"])]) + ")"
        for c in cartes)
    notes = ",\n ".join(f"({sql(c['id'])}::uuid, {sql(t)})" for c in cartes for t in c["notes"])
    trg = ", ".join(f"'{t}'" for t in sorted(TRIGGERS))
    with open(chemin, "w", encoding="utf-8", newline="\n") as f:
        f.write(f"""-- LOT PROD — import des cartes prospect Lauréats (Cégep de Saint-Hyacinthe × Football). Généré par
-- docs/import-laureats/outils/construire-lot-cartes.py — CONTIENT DES RENSEIGNEMENTS PERSONNELS : jamais dans git.
-- À envoyer TEL QUEL, en UN SEUL appel : une seule transaction. Une garde qui échoue annule tout.
-- {n} cartes, {nb_notes} notes. Insertion SOUS L'IDENTITÉ du recruteur football is_school_admin de l'unité (retrouvé
-- par requête) : cartes_prospect_avant_insert fixe l'unité et valide équipe/position ; le rapprochement s'enfile seul.
-- Aucun courriel, aucun téléphone → aucune invitation. Aucune fiche athletes touchée.
-- Rollback : docs/import-laureats/prod/rollback-cartes-laureats.sql (supprime exactement ces id).

create temp table _cp_ids (id uuid primary key);
insert into _cp_ids values
 {ids};

-- 1. Relevé avant (dans la transaction)
create temp table _cp_releve as
select txid_current() as txid,
       (select count(*) from public.cartes_prospect) as cartes,
       (select count(*) from public.cartes_prospect_notes) as notes,
       (select count(*) from public.cartes_prospect_invitations) as invitations,
       (select count(*) from public.athletes) as athletes,
       (select md5(coalesce(string_agg(a::text, '|' order by a.id), '')) from public.athletes a) as empreinte_athletes,
       (select count(*) from public.teams) as teams,
       (select count(*) from public.admin_operations where operation = 'IMPORT_CARTES_LAUREATS') as op_deja,
       (select id from public.sports where nom = 'Football') as football,
       (select count(*) from public.users u where u.school_id = '{UNITE_CEGEP}' and u.role = 'RECRUTEUR'::public.user_role
           and u.is_school_admin and u.sport_id = (select id from public.sports where nom = 'Football')) as nb_signataires,
       (select u.id from public.users u where u.school_id = '{UNITE_CEGEP}' and u.role = 'RECRUTEUR'::public.user_role
           and u.is_school_admin and u.sport_id = (select id from public.sports where nom = 'Football') limit 1) as signataire,
       (select id from public.users where email = 'bptds22@gmail.com') as par;

-- 2. Gardes AVANT toute écriture
do $g$
declare n int; t text[];
begin
  if (select txid from _cp_releve) <> txid_current() then
    raise exception 'NEXUS LOT: les instructions ne partagent pas une seule transaction — arrêt avant toute écriture';
  end if;
  if (select op_deja from _cp_releve) <> 0 then raise exception 'NEXUS LOT: import déjà appliqué (admin_operations)'; end if;
  if (select nb_signataires from _cp_releve) <> 1 then
    raise exception 'NEXUS LOT: % recruteur(s) football admin dans l''unité, 1 attendu', (select nb_signataires from _cp_releve);
  end if;
  if (select par from _cp_releve) is null then raise exception 'NEXUS LOT: compte admin introuvable'; end if;
  -- La RLS d'insertion (acces_carte_ecriture) exige user_has_pro() du signataire : abonnement pro/all_star actif.
  if not exists (select 1 from public.subscriptions s where s.user_id = (select signataire from _cp_releve)
                   and s.status = 'active' and s.tier in ('pro', 'all_star')) then
    raise exception 'NEXUS LOT: le signataire n''a pas d''abonnement Pro/All Star actif — la RLS refuserait les cartes';
  end if;
  if (select md5(pg_get_functiondef('public.cartes_prospect_avant_insert()'::regprocedure))) <> '{MD5_AVANT_INSERT}' then
    raise exception 'NEXUS LOT: cartes_prospect_avant_insert diffère de la définition vérifiée';
  end if;
  select array_agg(tgname::text order by tgname::text) into t from pg_trigger
   where tgrelid = 'public.cartes_prospect'::regclass and not tgisinternal;
  if t is distinct from array[{trg}]::text[] then
    raise exception 'NEXUS LOT: triggers de cartes_prospect = %, attendu les 12 vérifiés', t;
  end if;
  if (select count(*) from _cp_ids) <> {n} then raise exception 'NEXUS LOT: le plan ne compte pas {n} cartes'; end if;
  if exists (select 1 from public.cartes_prospect c join _cp_ids i on i.id = c.id) then
    raise exception 'NEXUS LOT: un des id prévus existe déjà';
  end if;
  select count(*) into n from public.teams t2
   where t2.id in ({", ".join(sql(e) + "::uuid" for e in equipes)})
     and t2.sport_id = (select football from _cp_releve) and t2.season = '2026-2027';
  if n <> {len(equipes)} then raise exception 'NEXUS LOT: % équipe(s) football 2026-2027 sur {len(equipes)}', n; end if;
end $g$;

-- 3. Insertions — sous l'identité du recruteur signataire (chemin de l'app)
select set_config('request.jwt.claims',
                  json_build_object('sub', (select signataire from _cp_releve), 'role', 'authenticated')::text, true);
set local role authenticated;

insert into public.cartes_prospect (id, team_id, prenom, nom, numero, position_id, promotion,
                                    taille_pieds, taille_pouces, poids_lbs, grade) values
 {lignes};

insert into public.cartes_prospect_notes (carte_id, contenu) values
 {notes};

reset role;
select set_config('request.jwt.claims', '', true);

insert into public.admin_operations (operation, motif, details, par)
select 'IMPORT_CARTES_LAUREATS',
       'Import de la liste de dépistage des Lauréats en cartes prospect (Cégep de Saint-Hyacinthe × Football) — GO BP',
       jsonb_build_object('carte_ids', (select jsonb_agg(id order by id) from _cp_ids), 'nb_cartes', {n},
                          'nb_notes', {nb_notes}, 'signataire', (select signataire from _cp_releve),
                          'rollback', 'docs/import-laureats/prod/rollback-cartes-laureats.sql'),
       (select par from _cp_releve);

-- 4. Gardes finales
do $f$
declare n int;
begin
  if (select txid from _cp_releve) <> txid_current() then raise exception 'NEXUS LOT: transaction changée'; end if;
  if (select count(*) from public.cartes_prospect) <> (select cartes from _cp_releve) + {n} then
    raise exception 'NEXUS LOT: cartes_prospect n''a pas grandi d''exactement {n} — tout est annulé';
  end if;
  if (select count(*) from public.cartes_prospect_notes) <> (select notes from _cp_releve) + {nb_notes} then
    raise exception 'NEXUS LOT: cartes_prospect_notes n''a pas grandi d''exactement {nb_notes} — tout est annulé';
  end if;
  select count(*) into n from public.cartes_prospect c join _cp_ids i on i.id = c.id
   where c.unite_cegep_id = '{UNITE_CEGEP}' and c.unite_sport_id = (select football from _cp_releve)
     and c.sport_athlete_id = (select football from _cp_releve) and c.cree_par = (select signataire from _cp_releve)
     and c.etape = 'IDENTIFIE' and c.promotion = {PROMOTION} and c.courriel is null and c.telephone is null
     and c.parent_courriel is null and c.invitation_etat is null and c.invitee_le is null and c.fusionnee_avec is null
     and c.school_id = (select t.school_id from public.teams t where t.id = c.team_id);
  if n <> {n} then raise exception 'NEXUS LOT: % carte(s) conforme(s) sur {n} — tout est annulé', n; end if;
  if (select count(*) from public.cartes_prospect_invitations) <> (select invitations from _cp_releve) then
    raise exception 'NEXUS LOT: une invitation a été créée — tout est annulé';
  end if;
  if (select count(*) from public.athletes) <> (select athletes from _cp_releve)
     or (select md5(coalesce(string_agg(a::text, '|' order by a.id), '')) from public.athletes a)
        <> (select empreinte_athletes from _cp_releve) then
    raise exception 'NEXUS LOT: une fiche athletes a changé — tout est annulé';
  end if;
  if (select count(*) from public.teams) <> (select teams from _cp_releve) then
    raise exception 'NEXUS LOT: teams a changé — tout est annulé';
  end if;
  if (select count(*) from public.admin_operations where operation = 'IMPORT_CARTES_LAUREATS') <> 1 then
    raise exception 'NEXUS LOT: ligne admin_operations absente ou en double — tout est annulé';
  end if;
end $f$;

-- 5. Résultat
select json_build_object(
  'avant', (select json_build_object('cartes', cartes, 'notes', notes, 'invitations', invitations, 'athletes', athletes) from _cp_releve),
  'apres', json_build_object('cartes', (select count(*) from public.cartes_prospect),
                             'notes', (select count(*) from public.cartes_prospect_notes),
                             'invitations', (select count(*) from public.cartes_prospect_invitations),
                             'athletes', (select count(*) from public.athletes)),
  'signataire', (select signataire from _cp_releve),
  'file_rapprochement', (select count(*) from public.rapprochement_file f join _cp_ids i on i.id = f.carte_id),
  'admin_operation', (select id from public.admin_operations where operation = 'IMPORT_CARTES_LAUREATS'),
  'meme_transaction', (select txid from _cp_releve) = txid_current()
) as resultat_lot_cartes;
""")
    return n, nb_notes


def ecrire_apercu(cartes, doublons, ignores, sans_nom, chemin):
    c = collections.Counter
    with open(chemin, "w", encoding="utf-8", newline="\n") as f:
        p = lambda *a: print(*a, file=f)
        p(f"APERÇU — import Lauréats (hors git). Onglets ignorés (cachés) : {ignores}")
        p(f"Cartes : {len(cartes)} | lignes sans nom ignorées : {sans_nom} | doublons même équipe écartés : {len(doublons)}")
        p(f"Notes : {sum(len(x['notes']) for x in cartes)} au total, sur {sum(1 for x in cartes if x['notes'])} cartes")
        p("\nGRADES :", dict(sorted(c(x["grade"] or "—" for x in cartes).items())))
        p("POSITIONS :", dict(c(x["position"] or "—" for x in cartes).most_common()))
        p(f"Numéro : {sum(1 for x in cartes if x['numero'])} | taille : {sum(1 for x in cartes if x['pieds'])} | "
          f"poids : {sum(1 for x in cartes if x['poids'])}")
        p("\nPAR ONGLET (équipe) :")
        for o, k in sorted(c(x["onglet"] for x in cartes).items()):
            p(f"  {o:28} {k:3}  {EQUIPES[o]}")
        p("\nDOUBLONS ÉCARTÉS (même équipe, 1re ligne gardée) :")
        for o, l, nm, l1 in doublons:
            p(f"  {o} L{l} « {nm} » — gardé : L{l1}")
        p("\nMÊME NOM DANS DEUX ÉCOLES (deux cartes) :")
        par = collections.defaultdict(list)
        for x in cartes:
            par[cle_nom(x["nom_complet"])].append(x["onglet"])
        for k, v in sorted(par.items()):
            if len(v) > 1:
                p(f"  {k} : {v}")
        p(f"\nNOTES « {NOTE_NOM} » :")
        for x in cartes:
            if NOTE_NOM in x["notes"] or NOTE_NOM_INCOMPLET in x["notes"]:
                p(f"  {x['onglet']} L{x['ligne']} : brut « {x['nom_complet']} » → prénom « {x['prenom']} », nom « {x['nom']} »")
        p("\nRANGS NON CONVERTIS :")
        for x in cartes:
            for t in x["notes"]:
                if t.startswith("Rang Excel"):
                    p(f"  {x['onglet']} L{x['ligne']} : {t}")
        p("\nPOSITIONS SANS CORRESPONDANCE (aucune position posée) :", dict(c(x["pos_brute"] for x in cartes if x["pos_brute"] and not x["position"])))
        p("\nTAILLE / POIDS EXTRAITS (extrait → résultat) :")
        for x in cartes:
            if x["extrait"] or x["poids"]:
                p(f"  {x['onglet']} L{x['ligne']} : « {x['extrait']} » → {x['pieds']}'{x['pouces']} / {x['poids']} lb")


def ecrire_rollback(cartes, chemin):
    """Ne contient que des id : versionnable. Refuse s'il y a déjà du travail accroché aux cartes."""
    n = len(cartes)
    ids = ",\n ".join(f"('{c['id']}'::uuid)" for c in sorted(cartes, key=lambda c: c["id"]))
    with open(chemin, "w", encoding="utf-8", newline="\n") as f:
        f.write(f"""-- ROLLBACK — import des cartes Lauréats (lot-cartes-laureats.sql). UN SEUL envoi, une seule transaction.
-- Généré par docs/import-laureats/outils/construire-lot-cartes.py. Ne contient que des id.
-- Supprime EXACTEMENT ces {n} cartes (notes, journal, file et propositions de rapprochement partent en cascade ;
-- chaque retrait est tracé dans cartes_prospect_suppressions). REFUSE si du travail s'y est accroché :
-- liaison décidée, fusion, invitation, rappel, liste, ou étape changée. À passer AVANT le rollback des équipes.

create temp table _rbc_ids (id uuid primary key);
insert into _rbc_ids values
 {ids};

create temp table _rbc_releve as
select txid_current() as txid, (select count(*) from public.cartes_prospect) as cartes,
       (select id from public.users where email = 'bptds22@gmail.com') as par;

do $g$
declare n int;
begin
  if (select count(*) from _rbc_ids) <> {n} then raise exception 'NEXUS ROLLBACK: la liste ne compte pas {n} id'; end if;
  if (select count(*) from public.cartes_prospect c join _rbc_ids i on i.id = c.id) <> {n} then
    raise exception 'NEXUS ROLLBACK: les {n} cartes ne sont pas toutes présentes — rien n''est supprimé';
  end if;
  select (select count(*) from public.rapprochements r join _rbc_ids i on i.id = r.carte_id where r.statut <> 'PROPOSEE')
       + (select count(*) from public.cartes_prospect c join _rbc_ids i on i.id = c.id
           where c.fusionnee_avec is not null or c.etape <> 'IDENTIFIE' or c.courriel is not null or c.telephone is not null)
       + (select count(*) from public.cartes_prospect_invitations x join _rbc_ids i on i.id = x.carte_id)
       + (select count(*) from public.cartes_prospect_rappels x join _rbc_ids i on i.id = x.carte_id)
       + (select count(*) from public.cartes_prospect_listes x join _rbc_ids i on i.id = x.carte_id)
    into n;
  if n <> 0 then raise exception 'NEXUS ROLLBACK: % élément(s) de travail accroché(s) aux cartes — rien n''est supprimé', n; end if;
end $g$;

delete from public.cartes_prospect where id in (select id from _rbc_ids);

insert into public.admin_operations (operation, motif, details, par)
select 'ROLLBACK_IMPORT_CARTES_LAUREATS', 'Retrait des cartes de l''import Lauréats',
       jsonb_build_object('nb_cartes', {n}, 'carte_ids', (select jsonb_agg(id order by id) from _rbc_ids)), (select par from _rbc_releve);

do $f$
begin
  if (select txid from _rbc_releve) <> txid_current() then raise exception 'NEXUS ROLLBACK: transaction changée'; end if;
  if (select count(*) from public.cartes_prospect) <> (select cartes from _rbc_releve) - {n} then
    raise exception 'NEXUS ROLLBACK: cartes_prospect n''a pas diminué d''exactement {n} — tout est annulé';
  end if;
end $f$;

select json_build_object('avant', (select cartes from _rbc_releve), 'apres', (select count(*) from public.cartes_prospect),
                         'restantes', (select count(*) from public.cartes_prospect c join _rbc_ids i on i.id = c.id)) as resultat_rollback_cartes;
""")


if __name__ == "__main__":
    cartes, doublons, ignores, sans_nom = lire(sys.argv[1])
    sortie = sys.argv[2]
    if len(sys.argv) > 3:
        ecrire_rollback(cartes, sys.argv[3])
    n, nb_notes = ecrire_lot(cartes, f"{sortie}/lot-cartes-laureats.sql")
    ecrire_apercu(cartes, doublons, ignores, sans_nom, f"{sortie}/apercu-cartes-laureats.txt")
    with open(f"{sortie}/ids-cartes-laureats.txt", "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(sorted(x["id"] for x in cartes)) + "\n")
    print(f"{n} cartes, {nb_notes} notes, {len(doublons)} doublons écartés, {sans_nom} lignes sans nom")
